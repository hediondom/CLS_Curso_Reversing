# -*- coding: utf-8 -*-
"""Extrae el texto de un .docx a la carpeta _trabajo/extraido, sin tocar el original."""
import os
import sys

from docx import Document

SALIDA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "extraido")


def extraer(path):
    doc = Document(path)
    partes = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            partes.append(t)
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip().replace("\n", " ") for c in fila.cells]
            if any(celdas):
                partes.append(" | ".join(celdas))
    return "\n".join(partes)


def main():
    for src in sys.argv[1:]:
        os.makedirs(SALIDA, exist_ok=True)
        base = os.path.splitext(os.path.basename(src))[0]
        dst = os.path.join(SALIDA, base + ".txt")
        texto = extraer(src)
        with open(dst, "w", encoding="utf-8") as f:
            f.write(texto)
        print(f"{len(texto)} chars -> {dst}")


if __name__ == "__main__":
    main()
